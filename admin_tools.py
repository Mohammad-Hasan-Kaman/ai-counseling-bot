# -*- coding: utf-8 -*-
"""
Admin panel tools module for the AI counseling bot
Provides: Excel export of user data and broadcast (mass) messaging
This module is standalone and does not modify any existing code.
"""
import io
import json
import asyncio
import logging

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import USER_RECORDS_DB

log = logging.getLogger(__name__)

BROADCAST_BATCH_DELAY = 0.05   # Delay between sends to avoid Bale server rate limits
BROADCAST_REPORT_EVERY = 20    # Send a progress report every N recipients


# ══════════════════════════════════════════════
#  Section 1: Excel export of user data
# ══════════════════════════════════════════════

_HEADER_FILL = PatternFill("solid", fgColor="2E5E4E")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

USERS_HEADERS = REQUESTS_HEADERS = [
    "ردیف", "نام و نام خانوادگی", "شماره تماس", "شناسه بله (Chat ID)",
    "جنسیت", "سن", "موضوع مشاوره (مشکل اصلی)",
    "سابقه مراجعه قبلی به روانشناس/روانپزشک", "توضیحات درمان قبلی",
    "انتظار یا رویکرد موردنظر از مشاوره", "ترجیح جنسیت مشاور",
    "شعبه انتخابی",
    "GHQ - جسمانی‌شدن (سوماتیک)", "GHQ - اضطراب و بی‌خوابی",
    "GHQ - اختلال کارکرد اجتماعی", "GHQ - افسردگی",
    "نتیجه نهایی GHQ (امتیاز کل)", "وضعیت GHQ",
    "مشاوران پیشنهادی شده", "دلیل پیشنهاد", "شماره مراجعه", "تاریخ ثبت درخواست",
]


def _style_header(ws, headers):
    fill, font, align = _HEADER_FILL, _HEADER_FONT, _HEADER_ALIGN
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.freeze_panes = "A2"
    for col_idx in range(1, len(headers) + 1):
        width = max(14, min(30, len(headers[col_idx - 1]) * 1.6 + 6))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _fmt_recs(recommendations_json: str | None) -> tuple[str, str]:
    """Extract consultant names and recommendation reasons from the stored JSON"""
    if not recommendations_json:
        return "", ""
    try:
        import json
        recs = json.loads(recommendations_json)
        if isinstance(recs, list):
            names = " | ".join(r.get("name", "") for r in recs if isinstance(r, dict))
            reasons = " | ".join(
                f"{r.get('name','')}: {r.get('reason','')}" for r in recs if isinstance(r, dict)
            )
            return names, reasons
        return str(recs), ""
    except Exception:
        return "", ""


def _ghq_status(total) -> str:
    """Status label based on the total GHQ score (severity threshold: 43)"""
    if total is None:
        return "انجام نداده"
    if total >= 43:
        return "شدید"
    if total >= 24:
        return "متوسط"
    return "طبیعی"


def export_users_excel() -> bytes:
    """
    Generate the user-data Excel output in memory (no file written to disk).
    Both sheets are complete and identical: sheet 1 "Users" (sorted by name) and
    sheet 2 "Request history" (sorted by date, newest first) —
    each row = one request with all information asked from the user.
    Returns: xlsx file bytes
    """
    import sqlite3
    conn = sqlite3.connect(str(USER_RECORDS_DB))

    requests_rows_name = conn.execute("""
        SELECT full_name, phone, telegram_id, gender, age, topic,
               has_prev_therapy, prev_detail, expectation, preferred_gender,
               branch, ghq_scores, ghq_total, recommendations, request_number, created_at
        FROM user_requests ORDER BY full_name COLLATE NOCASE ASC, request_number ASC
    """).fetchall()

    requests_rows_date = conn.execute("""
        SELECT full_name, phone, telegram_id, gender, age, topic,
               has_prev_therapy, prev_detail, expectation, preferred_gender,
               branch, ghq_scores, ghq_total, recommendations, request_number, created_at
        FROM user_requests ORDER BY created_at DESC
    """).fetchall()
    conn.close()

    wb = Workbook()

    def _write_request_row(ws, i, r):
        (name, phone, tg_id, gender, age, topic, has_prev, prev_detail,
         expectation, pref_gender, branch, ghq_json, ghq_total, recs,
         req_number, created) = r

        # Extract GHQ subscales from the stored JSON
        somatic = anxiety = social = depression = None
        if ghq_json:
            try:
                g = json.loads(ghq_json)
                somatic = g.get("somatic")
                anxiety = g.get("anxiety")
                social = g.get("social")
                depression = g.get("depression")
                if ghq_total is None and isinstance(g.get("total"), int):
                    ghq_total = g["total"]
            except Exception:
                pass

        rec_names, rec_reasons = _fmt_recs(recs)

        def _clean(v):
            """Convert None and the string 'None' to an empty string for clean display in Excel"""
            if v is None:
                return ""
            s = str(v)
            return "" if s.strip().lower() == "none" else s

        ws.append([
            i, name, phone, tg_id or "",
            gender or "",
            age if age is not None else "",
            topic or "",
            "بله" if has_prev else "خیر",
            _clean(prev_detail),
            expectation or "",
            pref_gender or "",
            branch or "-",
            somatic if somatic is not None else "-",
            anxiety if anxiety is not None else "-",
            social if social is not None else "-",
            depression if depression is not None else "-",
            ghq_total if ghq_total is not None else "-",
            _ghq_status(ghq_total),
            rec_names,
            rec_reasons,
            req_number if req_number is not None else "",
            created or "",
        ])

    # ── Sheet 1: Users (all information, sorted by name) ──
    ws_users = wb.active
    ws_users.title = "کاربران"
    _style_header(ws_users, USERS_HEADERS)
    for i, r in enumerate(requests_rows_name, start=1):
        _write_request_row(ws_users, i, r)

    # ── Sheet 2: Request history (same columns, sorted by date) ──
    ws_req = wb.create_sheet("سوابق درخواست‌ها")
    _style_header(ws_req, REQUESTS_HEADERS)
    for i, r in enumerate(requests_rows_date, start=1):
        _write_request_row(ws_req, i, r)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ══════════════════════════════════════════════
#  Section 2: Broadcast (mass messaging)
# ══════════════════════════════════════════════

def get_broadcast_targets() -> list[int]:
    """Get the IDs of all users who have interacted with the bot"""
    import sqlite3
    try:
        conn = sqlite3.connect(str(USER_RECORDS_DB))
        rows = conn.execute("SELECT DISTINCT first_telegram_id FROM users WHERE first_telegram_id IS NOT NULL").fetchall()
        conn.close()
        return [int(r[0]) for r in rows if r[0]]
    except Exception as e:
        log.error("Broadcast target fetch error: %s", e)
        return []


async def send_broadcast(bot, targets: list[int], from_chat_id: int, message_id: int, progress_msg=None) -> dict:
    """
    Copy an arbitrary message (text, photo, voice, file, ...) to all users.
    Uses copy_message so that any content type is supported.
    progress_msg: message whose progress report is updated (optional)
    """
    sent = failed = blocked = 0
    total = len(targets)

    async def _report():
        if progress_msg:
            text = (
                f"📤 **گزارش برادکست:**\n"
                f"✅ ارسال‌شده: {sent} / {total}\n"
                f"🚫 ناموفق (بلاک/غیرفعال): {failed}"
            )
            try:
                await progress_msg.edit_text(text, parse_mode="Markdown")
            except Exception:
                pass

    for idx, chat_id in enumerate(targets, start=1):
        try:
            await bot.copy_message(chat_id=chat_id, from_chat_id=from_chat_id, message_id=message_id)
            sent += 1
        except Exception as e:
            failed += 1
            msg_str = str(e).lower()
            if "blocked" in msg_str or "forbidden" in msg_str or "deactivated" in msg_str:
                blocked += 1
            log.debug("Broadcast fail to %s: %s", chat_id, e)

        if idx % BROADCAST_REPORT_EVERY == 0:
            await _report()
            await asyncio.sleep(BROADCAST_BATCH_DELAY * 10)

        await asyncio.sleep(BROADCAST_BATCH_DELAY)

    return {"sent": sent, "failed": failed, "blocked": blocked, "total": total}
